"""Stage 0 - get contract text and header facts out of an uploaded document."""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from datetime import date

from dateutil import parser as dateparser

from app.pipeline.segment import normalise

log = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".text"}

# \bgovern(ed|ing)\b, not a bare "govern" prefix - otherwise every mention of a
# "governmental entity" or "government agency" matches.
_GOVLAW_RE = re.compile(
    r"\bgovern(?:ed|ing)\b\s+(?:by|in\s+accordance\s+with)\s+"
    r"(?:and\s+construed\s+in\s+accordance\s+with\s+)?the\s+laws?\s+of\s+"
    r"(?:the\s+)?(?:State\s+of\s+|Commonwealth\s+of\s+)?"
    r"(?P<law>[A-Z][A-Za-z.' ]{2,44})",
    re.IGNORECASE,
)
_MONTH = (r"(?:January|February|March|April|May|June|July|August|September|"
          r"October|November|December)")
# Contracts write dates every way imaginable: "August 1, 2011", "1 August 2011",
# "1st day of August, 2011", "08/01/2011". Cover all four.
_EFFECTIVE_RE = re.compile(
    r"(?:Effective\s+Date|effective\s+as\s+of|dated\s+as\s+of|made\s+as\s+of|"
    r"entered\s+into\s+as\s+of|made\s+and\s+entered\s+into\s+(?:on|as\s+of)|"
    r"dated\s+this|dated)\s*[:,]?\s*"
    rf"(?P<date>{_MONTH}\s+\d{{1,2}}(?:st|nd|rd|th)?,?\s+\d{{4}}"
    rf"|\d{{1,2}}(?:st|nd|rd|th)?\s+(?:day\s+of\s+)?{_MONTH},?\s+\d{{4}}"
    r"|\d{1,2}/\d{1,2}/\d{2,4}|\d{4}-\d{2}-\d{2})",
    re.IGNORECASE,
)
_PARTIES_RE = re.compile(
    r"(?:by\s+and\s+between|between)\s+(?P<a>[A-Z][\w&.,'\- ]{3,70}?)\s+"
    r"(?:\([^)]{0,60}\)\s*)?(?:and|AND)\s+(?P<b>[A-Z][\w&.,'\- ]{3,70}?)\s*"
    r"(?:\(|,|\.|\n)",
)
_TYPE_HINTS: tuple[tuple[str, str], ...] = (
    ("distributor agreement", "Distribution"),
    ("distribution agreement", "Distribution"),
    ("license agreement", "License"),
    ("licensing agreement", "License"),
    ("supply agreement", "Supply"),
    ("services agreement", "Services"),
    ("service agreement", "Services"),
    ("consulting agreement", "Consulting"),
    ("employment agreement", "Employment"),
    ("non-disclosure", "NDA"),
    ("confidentiality agreement", "NDA"),
    ("joint venture", "Joint Venture"),
    ("reseller agreement", "Reseller"),
    ("franchise agreement", "Franchise"),
    ("maintenance agreement", "Maintenance"),
    ("development agreement", "Development"),
    ("manufacturing agreement", "Manufacturing"),
    ("co-branding", "Co-Branding"),
    ("marketing agreement", "Marketing"),
    ("agency agreement", "Agency"),
    ("hosting agreement", "Hosting"),
    ("sponsorship agreement", "Sponsorship"),
    ("outsourcing agreement", "Outsourcing"),
    ("transportation agreement", "Transportation"),
    ("promotion agreement", "Promotion"),
    ("endorsement agreement", "Endorsement"),
)


@dataclass
class IngestedDocument:
    title: str
    text: str
    contract_type: str | None = None
    party_a: str | None = None
    party_b: str | None = None
    effective_date: date | None = None
    governing_law: str | None = None
    pages: int | None = None


class UnsupportedDocument(ValueError):
    pass


# ------------------------------------------------------------------ readers
def _read_pdf(blob: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(blob))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception as exc:
            log.warning("pdf page extraction failed: %s", exc)
    return "\n\n".join(parts), len(reader.pages)


def _read_docx(blob: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(blob))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def read_document(filename: str, blob: bytes) -> tuple[str, int | None]:
    suffix = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if suffix not in SUPPORTED:
        raise UnsupportedDocument(
            f"unsupported file type '{suffix or filename}'. "
            f"Supported: {', '.join(sorted(SUPPORTED))}"
        )
    if suffix == ".pdf":
        return _read_pdf(blob)
    if suffix == ".docx":
        return _read_docx(blob), None
    for enc in ("utf-8", "latin-1"):
        try:
            return blob.decode(enc), None
        except UnicodeDecodeError:
            continue
    return blob.decode("utf-8", errors="replace"), None


# ------------------------------------------------------------- header facts
def guess_title(text: str, fallback: str) -> str:
    for line in text.split("\n")[:40]:
        s = line.strip()
        if 8 <= len(s) <= 140 and "agreement" in s.lower():
            return s[:400]
    for line in text.split("\n")[:15]:
        s = line.strip()
        if 8 <= len(s) <= 140:
            return s[:400]
    return fallback[:400]


def guess_contract_type(title: str, text: str) -> str | None:
    hay = (title + " " + text[:4000]).lower()
    for needle, label in _TYPE_HINTS:
        if needle in hay:
            return label
    return "Agreement" if "agreement" in hay else None


def guess_effective_date(text: str) -> date | None:
    head = text[:12000]
    for m in _EFFECTIVE_RE.finditer(head):
        raw = re.sub(r"(\d+)(st|nd|rd|th)\s+day\s+of\s+", r"\1 ", m.group("date"),
                     flags=re.IGNORECASE)
        try:
            d = dateparser.parse(raw, fuzzy=True).date()
        except (ValueError, OverflowError, TypeError):
            continue
        if 1980 <= d.year <= 2100:
            return d
    return None


def guess_parties(text: str) -> tuple[str | None, str | None]:
    m = _PARTIES_RE.search(text[:8000])
    if not m:
        return None, None

    def clean(s: str) -> str:
        s = re.sub(r"\s+", " ", s).strip(" ,.;:")
        return s[:250]

    return clean(m.group("a")) or None, clean(m.group("b")) or None


_GOVLAW_ADJ_RE = re.compile(
    r"\b(?:governed\s+by|subject\s+to|construed\s+in\s+accordance\s+with)\s+"
    r"(?:the\s+)?(?P<law>[A-Z][A-Za-z]{2,24})\s+law\b",
)


def guess_governing_law(text: str) -> str | None:
    m = _GOVLAW_RE.search(text) or _GOVLAW_ADJ_RE.search(text)
    if not m:
        return None
    from app.eval.cuad import trim_jurisdiction

    law = re.sub(r"\s+", " ", m.group("law")).strip(" ,.;:")
    law = re.split(r"\b(?:without|excluding|applicable|in\s+all|as\s+to)\b",
                   law, flags=re.IGNORECASE)[0].strip(" ,.;:")
    return trim_jurisdiction(law)


def ingest(filename: str, blob: bytes) -> IngestedDocument:
    raw, pages = read_document(filename, blob)
    text = normalise(raw)
    if len(text.strip()) < 200:
        raise UnsupportedDocument(
            "could not extract readable text - the file may be a scanned image "
            "with no text layer"
        )
    title = guess_title(text, filename)
    party_a, party_b = guess_parties(text)
    return IngestedDocument(
        title=title,
        text=text,
        contract_type=guess_contract_type(title, text),
        party_a=party_a,
        party_b=party_b,
        effective_date=guess_effective_date(text),
        governing_law=guess_governing_law(text),
        pages=pages,
    )


def ingest_text(title: str, text: str, contract_type: str | None = None) -> IngestedDocument:
    text = normalise(text)
    party_a, party_b = guess_parties(text)
    return IngestedDocument(
        title=title[:400],
        text=text,
        contract_type=contract_type or guess_contract_type(title, text),
        party_a=party_a,
        party_b=party_b,
        effective_date=guess_effective_date(text),
        governing_law=guess_governing_law(text),
    )
